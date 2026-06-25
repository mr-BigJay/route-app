from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from PySide6.QtCore import Qt
from PySide6.QtGui import QTextDocument
from PySide6.QtPrintSupport import QPrinter
from PySide6.QtWidgets import (
    QFileDialog,
    QFrame,
    QGridLayout,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QMessageBox,
    QPushButton,
    QScrollArea,
    QSizePolicy,
    QStackedWidget,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from database.db import DatabaseError, DatabaseManager
from ui.mission_form_widget import MissionFormWidget
from ui.utils import (
    Page,
    confirm,
    gregorian_to_jalali,
    make_stat_card,
    show_error,
    show_success,
    to_persian_digits,
)


class MissionsPage(Page):
    HOME_VIEW = 0
    FORM_VIEW = 1
    LIST_VIEW = 2

    def __init__(self, db: DatabaseManager) -> None:
        super().__init__("ماموریت‌ها", "ثبت و مدیریت ماموریت خودروهای سازمانی")
        self.db = db
        self.missions_cache: list[dict] = []

        self.stack = QStackedWidget()
        self.stack.addWidget(self._build_home_view())
        self.stack.addWidget(self._build_form_view())
        self.stack.addWidget(self._build_list_view())
        self.root_layout.addWidget(self.stack, stretch=1)

    def _build_home_view(self) -> QWidget:
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(16)

        self.stats_layout = QGridLayout()
        self.stats_layout.setSpacing(12)
        layout.addLayout(self.stats_layout)

        options_card = self.card()
        options_layout = QHBoxLayout(options_card)
        options_layout.setContentsMargins(24, 24, 24, 24)
        options_layout.setSpacing(18)
        new_button = self._option_button("ثبت ماموریت جدید", "success")
        list_button = self._option_button("لیست ماموریت‌ها", "default")
        new_button.clicked.connect(self.open_new_mission_form)
        list_button.clicked.connect(self.open_missions_list)
        options_layout.addWidget(new_button)
        options_layout.addWidget(list_button)
        layout.addWidget(options_card)

        today_card = self.card()
        today_layout = QVBoxLayout(today_card)
        today_layout.setContentsMargins(18, 16, 18, 16)
        title = QLabel("لیست ماموریت‌های امروز")
        title.setObjectName("sectionTitle")
        self.today_table = QTableWidget(0, 7)
        self.today_table.setHorizontalHeaderLabels(
            ["ردیف", "ساعت", "راننده", "خودرو", "مبدا", "مقصد", "مسافت"]
        )
        self._setup_table(self.today_table)
        today_layout.addWidget(title)
        today_layout.addWidget(self.today_table)
        layout.addWidget(today_card, stretch=1)
        return page

    def _option_button(self, title: str, variant: str) -> QPushButton:
        button = QPushButton(title)
        button.setObjectName("optionButton")
        button.setProperty("variant", variant)
        button.setCursor(Qt.CursorShape.PointingHandCursor)
        button.setMinimumHeight(90)
        return button

    def _build_form_view(self) -> QWidget:
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        page = QWidget()
        page.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        layout = QVBoxLayout(page)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(8)

        card = self.card()
        card.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        form_layout = QVBoxLayout(card)
        form_layout.setContentsMargins(14, 12, 14, 12)
        form_layout.setSpacing(8)
        self.mission_form = MissionFormWidget(self.db)
        self.mission_form.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred)
        self.mission_form.saved.connect(self.back_to_home)
        self.mission_form.cancelled.connect(self.back_to_home)
        form_layout.addWidget(self.mission_form, 1)
        layout.addWidget(card)
        layout.addStretch(1)
        scroll.setWidget(page)
        return scroll

    def _build_list_view(self) -> QWidget:
        card = self.card()
        layout = QVBoxLayout(card)
        layout.setContentsMargins(18, 16, 18, 16)
        layout.setSpacing(12)

        header = QHBoxLayout()
        title = QLabel("لیست ماموریت‌ها")
        title.setObjectName("sectionTitle")
        back_button = self.action_button("بازگشت", "ghost")
        back_button.clicked.connect(self.back_to_home)
        header.addWidget(title)
        header.addStretch(1)
        header.addWidget(back_button)

        toolbar = QHBoxLayout()
        report_button = self.action_button("گزارش‌گیری", "secondary")
        edit_button = self.action_button("ویرایش", "secondary")
        delete_button = self.action_button("حذف", "danger")
        report_button.clicked.connect(self.report_missions)
        edit_button.clicked.connect(self.edit_selected_mission)
        delete_button.clicked.connect(self.delete_selected_mission)
        toolbar.addWidget(report_button)
        toolbar.addWidget(edit_button)
        toolbar.addWidget(delete_button)
        toolbar.addStretch(1)

        self.list_table = QTableWidget(0, 9)
        self.list_table.setHorizontalHeaderLabels(
            ["شناسه", "تاریخ", "ساعت", "راننده", "خودرو", "مبدا", "مقصد", "مسافت", "توضیحات"]
        )
        self._setup_table(self.list_table)

        layout.addLayout(header)
        layout.addLayout(toolbar)
        layout.addWidget(self.list_table)
        return card

    def _setup_table(self, table: QTableWidget) -> None:
        table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        table.verticalHeader().setVisible(False)
        table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)

    def refresh(self) -> None:
        self._refresh_home()
        if hasattr(self, "mission_form"):
            self.mission_form.refresh_combos()
        if hasattr(self, "list_table"):
            self._refresh_list_table()

    def _refresh_home(self) -> None:
        self._refresh_stats()
        self._refresh_today_table()

    def _refresh_stats(self) -> None:
        while self.stats_layout.count():
            item = self.stats_layout.takeAt(0)
            widget = item.widget()
            if widget:
                widget.deleteLater()
        stats = self._mission_period_stats()
        cards = [
            ("ماموریت‌های امروز", stats["today"], "#2563EB"),
            ("این هفته", stats["week"], "#7C3AED"),
            ("این ماه", stats["month"], "#F97316"),
            ("امسال", stats["year"], "#22C55E"),
            ("همه ماموریت‌ها", stats["all"], "#0F172A"),
        ]
        for col, (title, value, color) in enumerate(cards):
            self.stats_layout.addWidget(make_stat_card(title, str(value), color), 0, col)

    def _refresh_today_table(self) -> None:
        today = gregorian_to_jalali()
        rows = self.db.list_missions(filters={"mission_date": today})
        self._populate_table(
            self.today_table,
            rows,
            include_id=False,
            columns=["row", "mission_time", "driver_name", "vehicle", "origin", "destination", "distance"],
        )

    def _refresh_list_table(self) -> None:
        self.missions_cache = self.db.list_missions()
        self._populate_table(
            self.list_table,
            self.missions_cache,
            include_id=True,
            columns=[
                "id",
                "mission_date",
                "mission_time",
                "driver_name",
                "vehicle",
                "origin",
                "destination",
                "distance",
                "description",
            ],
        )

    def _populate_table(
        self,
        table: QTableWidget,
        rows: list[dict],
        include_id: bool,
        columns: list[str],
    ) -> None:
        table.setRowCount(len(rows))
        for row_index, mission in enumerate(rows):
            values: list[str] = []
            for column in columns:
                if column == "row":
                    values.append(to_persian_digits(row_index + 1))
                elif column == "distance":
                    values.append(to_persian_digits(f"{float(mission['distance'] or 0):.1f}"))
                else:
                    values.append(to_persian_digits(mission.get(column) or ""))
            for col, value in enumerate(values):
                item = QTableWidgetItem(value)
                item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                table.setItem(row_index, col, item)

    def open_new_mission_form(self) -> None:
        self.mission_form.prepare_new()
        self.stack.setCurrentIndex(self.FORM_VIEW)

    def open_missions_list(self) -> None:
        self._refresh_list_table()
        self.stack.setCurrentIndex(self.LIST_VIEW)

    def back_to_home(self) -> None:
        self.mission_form.selected_id = None
        self.mission_form.clear_form()
        self._refresh_home()
        self.stack.setCurrentIndex(self.HOME_VIEW)

    def edit_selected_mission(self) -> None:
        mission = self._current_list_mission()
        if not mission:
            show_error(self, "ابتدا یک ماموریت را انتخاب کنید.")
            return
        self.mission_form.load_mission(mission)
        self.stack.setCurrentIndex(self.FORM_VIEW)

    def delete_selected_mission(self) -> None:
        mission = self._current_list_mission()
        if not mission:
            show_error(self, "ابتدا یک ماموریت را انتخاب کنید.")
            return
        if not confirm(self, "آیا از حذف ماموریت انتخاب‌شده مطمئن هستید؟"):
            return
        try:
            self.db.delete_mission(int(mission["id"]))
            show_success(self, "ماموریت حذف شد.")
            self._refresh_list_table()
            self._refresh_home()
        except DatabaseError as exc:
            show_error(self, f"حذف ماموریت انجام نشد: {exc}")

    def report_missions(self) -> None:
        period_box = QMessageBox(self)
        period_box.setWindowTitle("نوع گزارش")
        period_box.setText("کدام گزارش را می‌خواهید؟")
        today_button = period_box.addButton("گزارش امروز", QMessageBox.ButtonRole.AcceptRole)
        week_button = period_box.addButton("گزارش هفته گذشته", QMessageBox.ButtonRole.AcceptRole)
        month_button = period_box.addButton("گزارش ماه گذشته", QMessageBox.ButtonRole.AcceptRole)
        period_box.addButton("انصراف", QMessageBox.ButtonRole.RejectRole)
        period_box.exec()
        clicked_period = period_box.clickedButton()
        if clicked_period == today_button:
            rows = self._missions_for_last_days(1)
            title = "گزارش امروز"
        elif clicked_period == week_button:
            rows = self._missions_for_last_days(7)
            title = "گزارش هفته گذشته"
        elif clicked_period == month_button:
            rows = self._missions_for_last_days(30)
            title = "گزارش ماه گذشته"
        else:
            return
        if not rows:
            show_error(self, "داده‌ای برای گزارش انتخاب‌شده وجود ندارد.")
            return

        format_box = QMessageBox(self)
        format_box.setWindowTitle("فرمت خروجی")
        format_box.setText("فرمت خروجی را انتخاب کنید.")
        excel_button = format_box.addButton("Excel", QMessageBox.ButtonRole.AcceptRole)
        pdf_button = format_box.addButton("PDF", QMessageBox.ButtonRole.AcceptRole)
        format_box.addButton("انصراف", QMessageBox.ButtonRole.RejectRole)
        format_box.exec()
        clicked_format = format_box.clickedButton()
        if clicked_format == excel_button:
            self.export_excel(rows, title)
        elif clicked_format == pdf_button:
            self.export_pdf(rows, title)

    def export_excel(self, rows: list[dict], title: str) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            show_error(self, "کتابخانه openpyxl نصب نیست.")
            return
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "ذخیره گزارش ماموریت",
            str(Path.home() / "route-missions-report.xlsx"),
            "Excel Files (*.xlsx)",
        )
        if not file_path:
            return
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Missions"
        sheet.sheet_view.rightToLeft = True
        headers = ["ردیف", "تاریخ", "ساعت", "راننده", "خودرو", "مبدا", "مقصد", "مسافت", "سرنشینان", "توضیحات"]
        sheet.append([title])
        sheet.append(headers)
        for index, mission in enumerate(rows, start=1):
            sheet.append(
                [
                    index,
                    to_persian_digits(mission["mission_date"]),
                    to_persian_digits(mission["mission_time"]),
                    mission["driver_name"],
                    mission["vehicle"],
                    mission["origin"],
                    mission["destination"],
                    to_persian_digits(f"{float(mission['distance'] or 0):.1f}"),
                    mission["passengers"] or "",
                    mission["description"] or "",
                ]
            )
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        sheet.cell(1, 1).font = Font(bold=True, size=14)
        sheet.cell(1, 1).alignment = Alignment(horizontal="center")
        header_fill = PatternFill("solid", fgColor="2563EB")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[2]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for row in sheet.iter_rows(min_row=3):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[column[0].column_letter].width = min(max_length + 4, 48)
        workbook.save(file_path)
        show_success(self, "گزارش Excel با موفقیت ذخیره شد.")

    def export_pdf(self, rows: list[dict], title: str) -> None:
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "ذخیره گزارش PDF",
            str(Path.home() / "route-missions-report.pdf"),
            "PDF Files (*.pdf)",
        )
        if not file_path:
            return
        html_rows = ""
        for index, mission in enumerate(rows, start=1):
            distance_text = to_persian_digits(f"{float(mission['distance'] or 0):.1f}")
            html_rows += (
                "<tr>"
                f"<td>{to_persian_digits(index)}</td><td>{to_persian_digits(mission['mission_date'])}</td><td>{to_persian_digits(mission['mission_time'])}</td>"
                f"<td>{mission['driver_name']}</td><td>{mission['vehicle']}</td>"
                f"<td>{mission['origin']}</td><td>{mission['destination']}</td>"
                f"<td>{distance_text}</td>"
                "</tr>"
            )
        html = f"""
        <html dir="rtl">
        <head>
        <meta charset="utf-8">
        <style>
            body {{ font-family: Tahoma; direction: rtl; }}
            table {{ border-collapse: collapse; width: 100%; }}
            th, td {{ border: 1px solid #cccccc; padding: 6px; text-align: center; }}
            th {{ background: #2563EB; color: white; }}
        </style>
        </head>
        <body>
            <h2>{title}</h2>
            <table>
                <tr>
                    <th>ردیف</th><th>تاریخ</th><th>ساعت</th><th>راننده</th>
                    <th>خودرو</th><th>مبدا</th><th>مقصد</th><th>مسافت</th>
                </tr>
                {html_rows}
            </table>
        </body>
        </html>
        """
        document = QTextDocument()
        document.setHtml(html)
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(file_path)
        document.print_(printer)
        show_success(self, "گزارش PDF با موفقیت ذخیره شد.")

    def _current_list_mission(self) -> dict | None:
        row = self.list_table.currentRow()
        if row < 0 or row >= len(self.missions_cache):
            return None
        return self.missions_cache[row]

    def _mission_period_stats(self) -> dict[str, int]:
        rows = self.db.list_missions()
        today = gregorian_to_jalali()
        week_dates = set(self._jalali_dates_for_last_days(7))
        month_prefix = today[:7]
        year_prefix = today[:4]
        return {
            "today": sum(1 for row in rows if row["mission_date"] == today),
            "week": sum(1 for row in rows if row["mission_date"] in week_dates),
            "month": sum(1 for row in rows if row["mission_date"].startswith(month_prefix)),
            "year": sum(1 for row in rows if row["mission_date"].startswith(year_prefix)),
            "all": len(rows),
        }

    def _missions_for_last_days(self, days: int) -> list[dict]:
        dates = set(self._jalali_dates_for_last_days(days))
        return [mission for mission in self.db.list_missions() if mission["mission_date"] in dates]

    @staticmethod
    def _jalali_dates_for_last_days(days: int) -> list[str]:
        return [gregorian_to_jalali(date.today() - timedelta(days=offset)) for offset in range(days)]
