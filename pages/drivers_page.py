from __future__ import annotations

import re
from pathlib import Path

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QAbstractItemView,
    QCheckBox,
    QFileDialog,
    QFrame,
    QGraphicsBlurEffect,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPushButton,
    QScrollArea,
    QSizePolicy,
    QStackedWidget,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from database.db import DatabaseError, DatabaseManager
from ui.driver_status_badge import DriverStatusBadge
from ui.form_widgets import NoWheelComboBox, configure_combo_field, configure_line_edit_field
from ui.utils import Page, confirm, show_error, show_success, to_english_digits, to_persian_digits


PERSIAN_TEXT_RE = re.compile(r"^[\u0600-\u06FF\s‌]+$")
DATE_RE = re.compile(r"^\d{4}/\d{2}/\d{2}$")
VEHICLE_STATUS_GOVERNMENT = "دولتی"
VEHICLE_STATUS_RENTAL = "استیجاری"


class DriversPage(Page):
    OPTION_VIEW = 0
    FORM_VIEW = 1
    LIST_VIEW = 2

    def __init__(self, db: DatabaseManager) -> None:
        super().__init__("مدیریت رانندگان", "ثبت راننده جدید و مدیریت لیست رانندگان")
        self.setObjectName("driversPage")
        self.db = db
        self.drivers_cache: list[dict] = []
        self.editing_driver_id: int | None = None
        self.report_selection_mode = False
        self._updating_checks = False
        self._formatting_birth_date = False
        self._formatting_distance_rate = False
        self._drivers_page = 0
        self._drivers_page_size = 10

        self.drivers_summary_card = self._build_drivers_summary()
        self.root_layout.addWidget(self.drivers_summary_card, alignment=Qt.AlignmentFlag.AlignHCenter)

        self.stack = QStackedWidget()
        self.stack.addWidget(self._build_options_view())
        self.stack.addWidget(self._build_form_view())
        self.stack.addWidget(self._build_list_view())
        self.root_layout.addWidget(self.stack, stretch=1)
        self._build_profile_overlay()
        self._refresh_active_count()
        self._center_page_header()

    def _center_page_header(self) -> None:
        header = self.root_layout.itemAt(0).widget()
        if header is None:
            return
        layout = header.layout()
        if layout is None:
            return
        for index in range(layout.count()):
            item = layout.itemAt(index)
            widget = item.widget()
            if widget is not None:
                widget.setAlignment(Qt.AlignmentFlag.AlignCenter)
                layout.setAlignment(widget, Qt.AlignmentFlag.AlignHCenter)

    def _build_drivers_summary(self) -> QFrame:
        card = QFrame()
        card.setObjectName("driversSummaryCard")
        layout = QHBoxLayout(card)
        layout.setContentsMargins(28, 18, 28, 18)
        layout.setSpacing(36)

        self.total_count_label = QLabel(to_persian_digits(0))
        self.total_count_label.setObjectName("driversSummaryValue")
        self.active_count_label = QLabel(to_persian_digits(0))
        self.active_count_label.setObjectName("driversSummaryValueActive")
        self.inactive_count_label = QLabel(to_persian_digits(0))
        self.inactive_count_label.setObjectName("driversSummaryValueInactive")

        for title, value_label in [
            ("کل رانندگان", self.total_count_label),
            ("رانندگان فعال", self.active_count_label),
            ("رانندگان غیرفعال", self.inactive_count_label),
        ]:
            layout.addWidget(self._summary_stat_block(title, value_label))

        return card

    def _summary_stat_block(self, title: str, value_label: QLabel) -> QWidget:
        block = QWidget()
        block_layout = QVBoxLayout(block)
        block_layout.setContentsMargins(0, 0, 0, 0)
        block_layout.setSpacing(6)
        title_label = QLabel(title)
        title_label.setObjectName("driversSummaryTitle")
        title_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        value_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        block_layout.addWidget(title_label)
        block_layout.addWidget(value_label)
        return block

    def _build_options_view(self) -> QFrame:
        card = self.card()
        layout = QVBoxLayout(card)
        layout.setContentsMargins(26, 24, 26, 24)
        layout.setSpacing(18)

        title = QLabel("لطفاً یکی از گزینه‌های زیر را انتخاب کنید")
        title.setObjectName("sectionTitle")
        layout.addWidget(title, alignment=Qt.AlignmentFlag.AlignRight)

        buttons_layout = QHBoxLayout()
        buttons_layout.setSpacing(18)
        new_driver = self._option_button("ثبت راننده جدید", "success")
        drivers_list = self._option_button("لیست رانندگان", "default")
        new_driver.clicked.connect(self.open_new_driver_form)
        drivers_list.clicked.connect(self.open_drivers_list)
        buttons_layout.addWidget(new_driver)
        buttons_layout.addWidget(drivers_list)
        layout.addLayout(buttons_layout)
        layout.addStretch(1)
        return card

    def _option_button(self, title: str, variant: str) -> QPushButton:
        button = QPushButton(title)
        button.setObjectName("optionButton")
        button.setProperty("variant", variant)
        button.setCursor(Qt.CursorShape.PointingHandCursor)
        button.setMinimumHeight(92)
        return button

    def _build_form_view(self) -> QFrame:
        card = self.card()
        card.setObjectName("driverFormCard")
        layout = QVBoxLayout(card)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)

        title_bar = QFrame()
        title_bar.setObjectName("driverFormTitleBar")
        title_layout = QHBoxLayout(title_bar)
        title_layout.setContentsMargins(18, 14, 18, 14)
        self.form_title = QLabel("ثبت راننده جدید")
        self.form_title.setObjectName("driverFormTitle")
        self.form_title.setAlignment(Qt.AlignmentFlag.AlignCenter)
        title_layout.addStretch(1)
        title_layout.addWidget(self.form_title)
        title_layout.addStretch(1)

        form_body = QWidget()
        form_body.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        body_layout = QVBoxLayout(form_body)
        body_layout.setContentsMargins(20, 18, 20, 18)
        body_layout.setSpacing(14)

        personal_group = QFrame()
        personal_group.setObjectName("driverFormGroup")
        personal_group.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        personal_layout = QVBoxLayout(personal_group)
        personal_layout.setContentsMargins(12, 14, 12, 12)
        personal_layout.setSpacing(10)
        personal_layout.addWidget(self._right_label_row("اطلاعات فردی", "driverSectionTitle"))

        self.first_name_input = QLineEdit()
        self.first_name_input.setPlaceholderText("فقط حروف فارسی")
        self.last_name_input = QLineEdit()
        self.last_name_input.setPlaceholderText("فقط حروف فارسی")
        self.mobile_input = QLineEdit()
        self.mobile_input.setPlaceholderText("مثال: 09123456789")
        self.mobile_input.setMaxLength(11)
        self.national_id_input = QLineEdit()
        self.national_id_input.setPlaceholderText("۱۰ رقم")
        self.national_id_input.setMaxLength(10)
        self.birth_date_input = QLineEdit()
        self.birth_date_input.setPlaceholderText("yyyy/mm/dd")
        self.birth_date_input.setMaxLength(10)
        self.birth_date_input.textEdited.connect(self._format_birth_date)
        for widget in [
            self.first_name_input,
            self.last_name_input,
            self.mobile_input,
            self.national_id_input,
            self.birth_date_input,
        ]:
            self._prepare_input(widget)

        personal_layout.addWidget(
            self._two_field_row("نام *", self.first_name_input, "نام خانوادگی *", self.last_name_input)
        )
        personal_layout.addWidget(
            self._two_field_row(
                "شماره موبایل *",
                self.mobile_input,
                "شماره ملی *",
                self.national_id_input,
            )
        )
        personal_layout.addWidget(self._field_box("تاریخ تولد *", self.birth_date_input))

        vehicle_group = QFrame()
        vehicle_group.setObjectName("driverFormGroup")
        vehicle_group.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        vehicle_layout = QVBoxLayout(vehicle_group)
        vehicle_layout.setContentsMargins(12, 14, 12, 12)
        vehicle_layout.setSpacing(10)
        vehicle_layout.addWidget(self._right_label_row("اطلاعات خودرو", "driverSectionTitle"))

        self.car_model_input = QLineEdit()
        self.car_model_input.setPlaceholderText("مثال: سمند")
        self.car_year_input = QLineEdit()
        self.car_year_input.setPlaceholderText("۴ رقم")
        self.car_year_input.setMaxLength(4)
        self.car_color_input = QLineEdit()
        self.car_color_input.setPlaceholderText("مثال: سفید")
        self.vehicle_status_combo = NoWheelComboBox()
        self.vehicle_status_combo.addItem(VEHICLE_STATUS_GOVERNMENT)
        self.vehicle_status_combo.addItem(VEHICLE_STATUS_RENTAL)
        self.vehicle_status_combo.currentTextChanged.connect(self._update_distance_rate_visibility)
        self.distance_rate_input = QLineEdit()
        self.distance_rate_input.setPlaceholderText("مثال: 230,000")
        self.distance_rate_input.textEdited.connect(self._format_distance_rate)
        for widget in [self.car_model_input, self.car_year_input, self.car_color_input, self.distance_rate_input]:
            self._prepare_input(widget)
        self._prepare_combo(self.vehicle_status_combo)

        vehicle_layout.addWidget(
            self._two_field_row("مدل خودرو *", self.car_model_input, "سال تولید *", self.car_year_input)
        )
        vehicle_layout.addWidget(
            self._two_field_row("رنگ خودرو *", self.car_color_input, "وضعیت خودرو *", self.vehicle_status_combo)
        )
        self.distance_rate_box = self._field_box("نرخ محاسبه به ریال *", self.distance_rate_input)
        vehicle_layout.addWidget(self.distance_rate_box)
        self._update_distance_rate_visibility()

        buttons = QHBoxLayout()
        buttons.setSpacing(10)
        self.save_button = self.action_button("ثبت")
        back_button = self.action_button("بازگشت", "ghost")
        self.save_button.clicked.connect(self.save_driver)
        back_button.clicked.connect(self.back_to_options)
        buttons.addStretch(1)
        buttons.addWidget(self.save_button)
        buttons.addWidget(back_button)
        buttons.addStretch(1)

        body_layout.addWidget(personal_group)
        body_layout.addWidget(vehicle_group)
        body_layout.addLayout(buttons)

        scroll = QScrollArea()
        scroll.setObjectName("driverFormScroll")
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.Shape.NoFrame)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        scroll.setWidget(form_body)

        layout.addWidget(title_bar)
        layout.addWidget(scroll)
        return card

    def _two_field_row(
        self,
        label_a: str,
        widget_a: QWidget,
        label_b: str,
        widget_b: QWidget,
        stretch_a: int = 1,
        stretch_b: int = 1,
    ) -> QWidget:
        row = QWidget()
        row.setLayoutDirection(Qt.LayoutDirection.LeftToRight)
        layout = QHBoxLayout(row)
        layout.setDirection(QHBoxLayout.Direction.LeftToRight)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(10)
        layout.addWidget(self._field_box(label_b, widget_b), stretch=stretch_b)
        layout.addWidget(self._field_box(label_a, widget_a), stretch=stretch_a)
        return row

    def _right_label_row(self, text: str, object_name: str) -> QWidget:
        row = QWidget()
        row.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        layout = QHBoxLayout(row)
        layout.setContentsMargins(0, 0, 0, 0)
        label = QLabel(text)
        label.setObjectName(object_name)
        label.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        label.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        layout.addWidget(label)
        layout.addStretch(1)
        return row

    def _field_box(self, label: str, widget: QWidget) -> QFrame:
        box = QFrame()
        box.setObjectName("driverFieldBox")
        box.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        layout = QVBoxLayout(box)
        layout.setContentsMargins(8, 6, 8, 8)
        layout.setSpacing(4)
        layout.addWidget(self._right_label_row(label, "fieldLabel"))
        widget.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        layout.addWidget(widget)
        return box

    def _prepare_input(self, widget: QLineEdit) -> None:
        widget.setObjectName("driverFormInput")
        widget.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        configure_line_edit_field(widget)
        self._ensure_input_alignment(widget)

    def _ensure_input_alignment(self, widget: QLineEdit) -> None:
        widget.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        widget.setAttribute(Qt.WidgetAttribute.WA_RightToLeft, True)
        placeholder = widget.placeholderText()
        if placeholder:
            widget.setPlaceholderText("")
            widget.setPlaceholderText(placeholder)
        widget.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)

    def _prepare_combo(self, combo: NoWheelComboBox) -> None:
        combo.setObjectName("driverFormInput")
        combo.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        combo.setAttribute(Qt.WidgetAttribute.WA_RightToLeft, True)
        combo.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        combo.setEditable(True)
        combo.currentTextChanged.connect(lambda _text: self._ensure_combo_alignment(combo))
        self._ensure_combo_alignment(combo)
        combo.view().setLayoutDirection(Qt.LayoutDirection.RightToLeft)

    def _ensure_combo_alignment(self, combo: NoWheelComboBox) -> None:
        line_edit = combo.lineEdit()
        if line_edit is None:
            return
        line_edit.setReadOnly(True)
        line_edit.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        line_edit.setAttribute(Qt.WidgetAttribute.WA_RightToLeft, True)
        configure_combo_field(combo)

    def _update_distance_rate_visibility(self) -> None:
        is_rental = self.vehicle_status_combo.currentText() == VEHICLE_STATUS_RENTAL
        self.distance_rate_box.setVisible(is_rental)
        if not is_rental:
            self.distance_rate_input.clear()

    def _driver_list_button(self, title: str, variant: str) -> QPushButton:
        button = QPushButton(title)
        button.setObjectName("driverListButton")
        button.setProperty("variant", variant)
        button.setCursor(Qt.CursorShape.PointingHandCursor)
        return button

    def _build_list_view(self) -> QFrame:
        card = self.card()
        card.setObjectName("missionsTableCard")
        layout = QVBoxLayout(card)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)

        header = QFrame()
        header.setObjectName("missionsTableHeader")
        header_layout = QVBoxLayout(header)
        header_layout.setContentsMargins(18, 16, 18, 14)
        header_layout.setSpacing(4)
        title = QLabel("لیست رانندگان")
        title.setObjectName("missionsTableTitle")
        title.setAlignment(Qt.AlignmentFlag.AlignCenter)
        title.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        subtitle = QLabel("مشاهده، ویرایش و مدیریت رانندگان ثبت‌شده")
        subtitle.setObjectName("missionsTableSubtitle")
        subtitle.setAlignment(Qt.AlignmentFlag.AlignCenter)
        subtitle.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        header_layout.addWidget(title)
        header_layout.addWidget(subtitle)

        wrap = QFrame()
        wrap.setObjectName("missionsTableWrap")
        body = QVBoxLayout(wrap)
        body.setContentsMargins(14, 14, 14, 14)
        body.setSpacing(12)

        toolbar = QHBoxLayout()
        toolbar.setSpacing(8)
        back_button = self.action_button("بازگشت", "ghost")
        back_button.clicked.connect(self.back_to_options)
        profile_button = self._driver_list_button("مشاهده پروفایل", "profile")
        edit_button = self._driver_list_button("ویرایش", "edit")
        delete_button = self.action_button("حذف", "danger")
        active_button = self._driver_list_button("فعالسازی", "activate")
        inactive_button = self._driver_list_button("غیرفعالسازی", "deactivate")
        self.report_button = self._driver_list_button("گزارش از لیست", "report")
        profile_button.clicked.connect(self.show_selected_driver_profile)
        edit_button.clicked.connect(self.edit_selected_driver)
        delete_button.clicked.connect(self.delete_selected_drivers)
        active_button.clicked.connect(self.activate_selected_drivers)
        inactive_button.clicked.connect(self.deactivate_selected_drivers)
        self.report_button.clicked.connect(self.report_from_list)
        toolbar.addWidget(back_button)
        toolbar.addStretch(1)
        for button in [
            profile_button,
            edit_button,
            delete_button,
            active_button,
            inactive_button,
            self.report_button,
        ]:
            toolbar.addWidget(button)

        self.table = QTableWidget(0, 6)
        self.table.setObjectName("missionsTable")
        self.table.setHorizontalHeaderLabels(
            [
                "انتخاب",
                "نام",
                "نام خانوادگی",
                "شماره موبایل",
                "مدل ماشین",
                "وضعیت",
            ]
        )
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
        self.table.horizontalHeader().setDefaultAlignment(Qt.AlignmentFlag.AlignCenter)
        self.table.verticalHeader().setVisible(False)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.setShowGrid(False)
        self.table.setFocusPolicy(Qt.FocusPolicy.StrongFocus)

        pagination_wrap = QFrame()
        pagination_wrap.setObjectName("driversTableFooter")
        pagination_wrap.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        pagination = QHBoxLayout(pagination_wrap)
        pagination.setContentsMargins(4, 10, 4, 2)
        pagination.setSpacing(12)

        page_size_row = QHBoxLayout()
        page_size_row.setDirection(QHBoxLayout.Direction.LeftToRight)
        page_size_row.setSpacing(8)
        show_prefix_label = QLabel("نمایش")
        show_prefix_label.setObjectName("driversPageSizeLabel")
        show_suffix_label = QLabel("مورد در هر صفحه")
        show_suffix_label.setObjectName("driversPageSizeLabel")
        self.page_size_combo = NoWheelComboBox()
        self.page_size_combo.setObjectName("driversPageSizeCombo")
        self.page_size_combo.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        for size in (10, 20, 50):
            self.page_size_combo.addItem(to_persian_digits(size), size)
        self.page_size_combo.setCurrentIndex(0)
        self.page_size_combo.currentIndexChanged.connect(self._on_page_size_changed)
        page_size_row.addWidget(show_prefix_label)
        page_size_row.addWidget(self.page_size_combo)
        page_size_row.addWidget(show_suffix_label)

        nav_row = QHBoxLayout()
        nav_row.setDirection(QHBoxLayout.Direction.LeftToRight)
        nav_row.setSpacing(6)
        self.range_info_label = QLabel()
        self.range_info_label.setObjectName("driversPageRangeInfo")
        self.range_info_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.prev_page_button = QPushButton("‹")
        self.prev_page_button.setObjectName("driversPageNavButton")
        self.prev_page_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.prev_page_button.setFixedSize(32, 32)
        self.prev_page_button.clicked.connect(self._go_prev_page)
        self.page_number_label = QLabel(to_persian_digits(1))
        self.page_number_label.setObjectName("driversPageNumber")
        self.page_number_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.page_number_label.setFixedSize(32, 32)
        self.next_page_button = QPushButton("›")
        self.next_page_button.setObjectName("driversPageNavButton")
        self.next_page_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.next_page_button.setFixedSize(32, 32)
        self.next_page_button.clicked.connect(self._go_next_page)
        nav_row.addWidget(self.range_info_label)
        nav_row.addWidget(self.prev_page_button)
        nav_row.addWidget(self.page_number_label)
        nav_row.addWidget(self.next_page_button)

        pagination.addLayout(nav_row)
        pagination.addStretch(1)
        pagination.addLayout(page_size_row)

        layout.addWidget(header)
        body.addLayout(toolbar)
        body.addWidget(self.table, stretch=1)
        body.addWidget(pagination_wrap)
        layout.addWidget(wrap, stretch=1)
        return card

    def _build_profile_overlay(self) -> None:
        self.profile_overlay = QFrame(self)
        self.profile_overlay.setObjectName("profileOverlay")
        self.profile_overlay.hide()

        overlay_layout = QVBoxLayout(self.profile_overlay)
        overlay_layout.setContentsMargins(24, 24, 24, 24)
        overlay_layout.addStretch(1)

        profile_card = QFrame()
        profile_card.setObjectName("profileCard")
        profile_card.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)
        profile_card.setFixedWidth(560)
        card_layout = QVBoxLayout(profile_card)
        card_layout.setContentsMargins(22, 20, 22, 20)
        card_layout.setSpacing(12)

        profile_title = QLabel("پروفایل راننده")
        profile_title.setObjectName("sectionTitle")
        self.profile_info_label = QLabel()
        self.profile_info_label.setObjectName("profileInfo")
        self.profile_info_label.setWordWrap(True)
        self.profile_info_label.setAlignment(Qt.AlignmentFlag.AlignRight)
        close_button = self.action_button("بازگشت", "ghost")
        close_button.clicked.connect(self.close_profile_overlay)

        card_layout.addWidget(profile_title, alignment=Qt.AlignmentFlag.AlignRight)
        card_layout.addWidget(self.profile_info_label)
        card_layout.addWidget(close_button, alignment=Qt.AlignmentFlag.AlignLeft)

        overlay_layout.addWidget(profile_card, alignment=Qt.AlignmentFlag.AlignCenter)
        overlay_layout.addStretch(1)

    def refresh(self) -> None:
        self._refresh_active_count()
        if hasattr(self, "table"):
            self.refresh_table()

    def open_new_driver_form(self) -> None:
        self.editing_driver_id = None
        self.form_title.setText("ثبت راننده جدید")
        self.save_button.setText("ثبت")
        self.clear_form()
        self.stack.setCurrentIndex(self.FORM_VIEW)

    def open_drivers_list(self) -> None:
        self.refresh_table()
        self.stack.setCurrentIndex(self.LIST_VIEW)

    def back_to_options(self) -> None:
        self.editing_driver_id = None
        self._leave_report_selection_mode()
        self.clear_form()
        self.stack.setCurrentIndex(self.OPTION_VIEW)

    def refresh_table(self) -> None:
        self._refresh_active_count()
        self.drivers_cache = self.db.list_drivers()
        self._clamp_drivers_page()
        self._populate_table_page()

    def _drivers_total_pages(self) -> int:
        if not self.drivers_cache:
            return 1
        return max(1, (len(self.drivers_cache) + self._drivers_page_size - 1) // self._drivers_page_size)

    def _clamp_drivers_page(self) -> None:
        max_page = self._drivers_total_pages() - 1
        if self._drivers_page > max_page:
            self._drivers_page = max(0, max_page)

    def _paginated_drivers(self) -> list[dict]:
        start = self._drivers_page * self._drivers_page_size
        end = start + self._drivers_page_size
        return self.drivers_cache[start:end]

    def _populate_table_page(self) -> None:
        page_drivers = self._paginated_drivers()
        self._updating_checks = True
        self.table.setRowCount(len(page_drivers))
        for row, driver in enumerate(page_drivers):
            checkbox = QCheckBox()
            checkbox.setProperty("driver_id", int(driver["id"]))
            checkbox.stateChanged.connect(self.on_driver_checked)
            checkbox_holder = QFrame()
            checkbox_layout = QHBoxLayout(checkbox_holder)
            checkbox_layout.setContentsMargins(0, 0, 0, 0)
            checkbox_layout.addWidget(checkbox, alignment=Qt.AlignmentFlag.AlignCenter)
            self.table.setCellWidget(row, 0, checkbox_holder)

            values = [
                driver["first_name"],
                driver["last_name"],
                to_persian_digits(driver["mobile"]),
                driver["car_model"],
            ]
            for col, value in enumerate(values, start=1):
                item = QTableWidgetItem(str(value))
                item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                self.table.setItem(row, col, item)

            status_holder = QFrame()
            status_layout = QHBoxLayout(status_holder)
            status_layout.setContentsMargins(0, 0, 0, 0)
            status_layout.addWidget(
                DriverStatusBadge(bool(int(driver.get("is_active", 1)))),
                alignment=Qt.AlignmentFlag.AlignCenter,
            )
            self.table.setCellWidget(row, 5, status_holder)
        self._updating_checks = False
        self._update_pagination_controls()

    def _update_pagination_controls(self) -> None:
        total = len(self.drivers_cache)
        total_pages = self._drivers_total_pages()
        if total == 0:
            range_start = 0
            range_end = 0
        else:
            range_start = self._drivers_page * self._drivers_page_size + 1
            range_end = min((self._drivers_page + 1) * self._drivers_page_size, total)
        self.range_info_label.setText(
            f"{to_persian_digits(range_start)}-{to_persian_digits(range_end)} "
            f"از {to_persian_digits(total)} مورد"
        )
        self.page_number_label.setText(to_persian_digits(self._drivers_page + 1))
        self.prev_page_button.setEnabled(self._drivers_page > 0)
        self.next_page_button.setEnabled(self._drivers_page < total_pages - 1)

    def _on_page_size_changed(self, _index: int) -> None:
        page_size = self.page_size_combo.currentData()
        if page_size is None:
            return
        self._drivers_page_size = int(page_size)
        self._drivers_page = 0
        self._populate_table_page()

    def _go_prev_page(self) -> None:
        if self._drivers_page > 0:
            self._drivers_page -= 1
            self._populate_table_page()

    def _go_next_page(self) -> None:
        if self._drivers_page < self._drivers_total_pages() - 1:
            self._drivers_page += 1
            self._populate_table_page()

    def collect_form_data(self) -> dict | None:
        vehicle_status = self.vehicle_status_combo.currentText().strip()
        distance_rate = (
            self._plain_number(self.distance_rate_input.text())
            if vehicle_status == VEHICLE_STATUS_RENTAL
            else 0
        )
        data = {
            "first_name": self.first_name_input.text().strip(),
            "last_name": self.last_name_input.text().strip(),
            "mobile": self._normalize_digits(self.mobile_input.text().strip()),
            "national_id": self._normalize_digits(self.national_id_input.text().strip()),
            "birth_date": to_english_digits(self.birth_date_input.text().strip()),
            "car_model": self.car_model_input.text().strip(),
            "car_year": to_english_digits(self.car_year_input.text().strip()),
            "car_color": self.car_color_input.text().strip(),
            "vehicle_status": vehicle_status,
            "distance_rate": distance_rate,
        }

        required_fields = [
            "first_name",
            "last_name",
            "mobile",
            "national_id",
            "birth_date",
            "car_model",
            "car_year",
            "car_color",
            "vehicle_status",
        ]
        if any(not data[key] for key in required_fields):
            show_error(self, "همه فیلدهای فرم ثبت راننده الزامی است.")
            return None
        if not self._is_persian_text(data["first_name"]):
            show_error(self, "نام فقط باید شامل حروف فارسی باشد.")
            return None
        if not self._is_persian_text(data["last_name"]):
            show_error(self, "نام خانوادگی فقط باید شامل حروف فارسی باشد.")
            return None
        if not re.fullmatch(r"09\d{9}", data["mobile"]):
            show_error(self, "شماره موبایل باید با 09 شروع شود و 11 رقم باشد.")
            return None
        if not re.fullmatch(r"\d{10}", data["national_id"]):
            show_error(self, "شماره ملی باید دقیقاً 10 رقم باشد.")
            return None
        if not self._is_valid_date(data["birth_date"]):
            show_error(self, "تاریخ تولد باید با فرمت yyyy/mm/dd وارد شود.")
            return None
        if not re.fullmatch(r"\d{4}", data["car_year"]):
            show_error(self, "سال تولید ماشین باید دقیقاً 4 رقم باشد.")
            return None
        if vehicle_status == VEHICLE_STATUS_RENTAL and data["distance_rate"] <= 0:
            show_error(self, "برای خودرو استیجاری، نرخ محاسبه به ریال الزامی است.")
            return None

        if self.editing_driver_id is not None:
            current = self._driver_by_id(self.editing_driver_id)
            if current:
                data["is_active"] = int(current.get("is_active", 1))
                data["inactive_reason"] = current.get("inactive_reason", "")
        return data

    def save_driver(self) -> None:
        data = self.collect_form_data()
        if data is None:
            return
        full_name = f"{data['first_name']} {data['last_name']}"
        action_text = "ویرایش" if self.editing_driver_id is not None else "ثبت"
        if not confirm(self, f"آیا {action_text} راننده «{full_name}» تایید می‌شود؟"):
            return
        try:
            if self.editing_driver_id is None:
                self.db.add_driver(data)
                show_success(self, "راننده با موفقیت ثبت شد.")
            else:
                self.db.update_driver(self.editing_driver_id, data)
                show_success(self, "اطلاعات راننده ویرایش شد.")
            self.clear_form()
            self.open_drivers_list()
        except DatabaseError as exc:
            show_error(self, f"ذخیره راننده انجام نشد: {exc}")

    def edit_selected_driver(self) -> None:
        self._leave_report_selection_mode()
        selected = self.selected_drivers()
        if len(selected) != 1:
            show_error(self, "برای ویرایش، دقیقاً یک راننده را انتخاب کنید.")
            return
        driver = selected[0]
        self.editing_driver_id = int(driver["id"])
        self.form_title.setText("ویرایش راننده")
        self.save_button.setText("ذخیره تغییرات")
        self.first_name_input.setText(driver["first_name"])
        self.last_name_input.setText(driver["last_name"])
        self.mobile_input.setText(to_persian_digits(driver["mobile"]))
        self.national_id_input.setText(to_persian_digits(driver["national_id"]))
        self.birth_date_input.setText(to_persian_digits(driver["birth_date"]))
        self.car_model_input.setText(driver["car_model"])
        self.car_year_input.setText(to_persian_digits(driver["car_year"]))
        self.car_color_input.setText(driver["car_color"])
        status = driver.get("vehicle_status") or VEHICLE_STATUS_GOVERNMENT
        status_index = self.vehicle_status_combo.findText(status)
        if status_index >= 0:
            self.vehicle_status_combo.setCurrentIndex(status_index)
        self._update_distance_rate_visibility()
        if status == VEHICLE_STATUS_RENTAL:
            self.distance_rate_input.setText(self._format_number(driver["distance_rate"]))
        self.stack.setCurrentIndex(self.FORM_VIEW)

    def delete_selected_drivers(self) -> None:
        self._leave_report_selection_mode()
        selected = self.selected_drivers()
        if not selected:
            show_error(self, "ابتدا راننده‌های مورد نظر را انتخاب کنید.")
            return
        if not confirm(self, f"آیا حذف {len(selected)} راننده انتخاب‌شده تایید می‌شود؟"):
            return
        try:
            for driver in selected:
                self.db.delete_driver(int(driver["id"]))
            show_success(self, "راننده‌های انتخاب‌شده حذف شدند.")
            self.refresh_table()
        except DatabaseError as exc:
            show_error(self, f"حذف راننده انجام نشد: {exc}")

    def activate_selected_drivers(self) -> None:
        self._leave_report_selection_mode()
        selected = self.selected_drivers()
        if not selected:
            show_error(self, "ابتدا راننده‌های مورد نظر را انتخاب کنید.")
            return
        try:
            for driver in selected:
                self.db.set_driver_active(int(driver["id"]), True)
            show_success(self, "راننده‌های انتخاب‌شده فعال شدند.")
            self.refresh_table()
        except DatabaseError as exc:
            show_error(self, f"فعال‌سازی راننده انجام نشد: {exc}")

    def deactivate_selected_drivers(self) -> None:
        self._leave_report_selection_mode()
        selected = self.selected_drivers()
        if not selected:
            show_error(self, "ابتدا راننده‌های مورد نظر را انتخاب کنید.")
            return

        reason_box = QMessageBox(self)
        reason_box.setWindowTitle("علت غیرفعال‌سازی")
        reason_box.setText("علت غیرفعال‌سازی راننده چیست؟")
        settlement = reason_box.addButton("تسویه حساب کرده", QMessageBox.ButtonRole.AcceptRole)
        retired = reason_box.addButton("بازنشست شده", QMessageBox.ButtonRole.AcceptRole)
        reason_box.addButton("انصراف", QMessageBox.ButtonRole.RejectRole)
        reason_box.exec()
        clicked = reason_box.clickedButton()
        if clicked not in (settlement, retired):
            return
        reason = clicked.text()

        try:
            for driver in selected:
                self.db.set_driver_active(int(driver["id"]), False, reason)
            show_success(self, "راننده‌های انتخاب‌شده غیرفعال شدند.")
            self.refresh_table()
        except DatabaseError as exc:
            show_error(self, f"غیرفعال‌سازی راننده انجام نشد: {exc}")

    def report_from_list(self) -> None:
        if self.report_selection_mode:
            rows = self.selected_drivers()
            if not rows:
                show_error(self, "ابتدا چک‌باکس راننده‌های مورد نظر برای گزارش را انتخاب کنید.")
                return
            self.export_drivers_excel(rows)
            self._leave_report_selection_mode(clear_checks=True)
            return

        report_box = QMessageBox(self)
        report_box.setWindowTitle("گزارش از لیست")
        report_box.setText("نوع گزارش را انتخاب کنید.")
        all_button = report_box.addButton("انتخاب همه", QMessageBox.ButtonRole.AcceptRole)
        selected_button = report_box.addButton("انتخاب راننده", QMessageBox.ButtonRole.AcceptRole)
        report_box.addButton("انصراف", QMessageBox.ButtonRole.RejectRole)
        report_box.exec()
        clicked = report_box.clickedButton()
        if clicked == all_button:
            rows = self.drivers_cache
            self.export_drivers_excel(rows)
        elif clicked == selected_button:
            self.report_selection_mode = True
            self.clear_driver_checks()
            self.report_button.setText("ایجاد گزارش انتخاب‌شده")
            show_success(self, "اکنون راننده‌های مورد نظر را انتخاب کنید و دوباره دکمه گزارش را بزنید.")
        else:
            return

    def export_drivers_excel(self, drivers: list[dict]) -> None:
        if not drivers:
            show_error(self, "داده‌ای برای گزارش وجود ندارد.")
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            show_error(self, "کتابخانه openpyxl نصب نیست.")
            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "ذخیره گزارش رانندگان",
            str(Path.home() / "route-drivers-report.xlsx"),
            "Excel Files (*.xlsx)",
        )
        if not file_path:
            return

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Drivers"
        sheet.sheet_view.rightToLeft = True
        headers = [
            "نام",
            "نام خانوادگی",
            "شماره موبایل",
            "شماره ملی",
            "تاریخ تولد",
            "مدل ماشین",
            "سال تولید ماشین",
            "رنگ ماشین",
            "نرخ محاسبه",
            "وضعیت",
        ]
        sheet.append(headers)
        for driver in drivers:
            sheet.append(
                [
                    driver["first_name"],
                    driver["last_name"],
                    driver["mobile"],
                    driver["national_id"],
                    driver["birth_date"],
                    driver["car_model"],
                    driver["car_year"],
                    driver["car_color"],
                    self._format_rial(driver["distance_rate"]),
                    "فعال" if int(driver.get("is_active", 1)) else "غیرفعال",
                ]
            )

        header_fill = PatternFill("solid", fgColor="2563EB")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[column[0].column_letter].width = min(max_length + 4, 45)
        workbook.save(file_path)
        show_success(self, "گزارش Excel رانندگان با موفقیت ذخیره شد.")

    def selected_drivers(self) -> list[dict]:
        driver_map = {int(driver["id"]): driver for driver in self.drivers_cache}
        selected: list[dict] = []
        for row in range(self.table.rowCount()):
            checkbox = self._checkbox_at_row(row)
            if not checkbox or not checkbox.isChecked():
                continue
            driver_id = checkbox.property("driver_id")
            if driver_id is None:
                continue
            driver = driver_map.get(int(driver_id))
            if driver is not None:
                selected.append(driver)
        return selected

    def on_driver_checked(self, state: int) -> None:
        if self._updating_checks or state != Qt.CheckState.Checked.value:
            return
        sender = self.sender()
        if not isinstance(sender, QCheckBox):
            return
        checked_driver_id = sender.property("driver_id")
        if self.report_selection_mode:
            return
        self._updating_checks = True
        for row in range(self.table.rowCount()):
            checkbox = self._checkbox_at_row(row)
            if checkbox and checkbox.property("driver_id") != checked_driver_id:
                checkbox.setChecked(False)
        self._updating_checks = False

    def show_selected_driver_profile(self) -> None:
        self._leave_report_selection_mode()
        selected = self.selected_drivers()
        if len(selected) != 1:
            show_error(self, "برای مشاهده پروفایل، دقیقاً یک راننده را انتخاب کنید.")
            return
        driver = selected[0]
        status = "فعال" if int(driver.get("is_active", 1)) else "غیرفعال"
        inactive_reason = driver.get("inactive_reason") or "-"
        self.profile_info_label.setText(
            "\n".join(
                [
                    "اطلاعات فردی",
                    f"نام: {driver['first_name']}",
                    f"نام خانوادگی: {driver['last_name']}",
                    f"شماره موبایل: {to_persian_digits(driver['mobile'])}",
                    f"شماره ملی: {to_persian_digits(driver['national_id'])}",
                    f"تاریخ تولد: {to_persian_digits(driver['birth_date'])}",
                    "",
                    "اطلاعات خودرو",
                    f"مدل ماشین: {driver['car_model']}",
                    f"سال تولید ماشین: {to_persian_digits(driver['car_year'])}",
                    f"رنگ ماشین: {driver['car_color']}",
                    f"وضعیت خودرو: {driver.get('vehicle_status') or VEHICLE_STATUS_GOVERNMENT}",
                    f"نرخ محاسبه: {self._format_rial(driver['distance_rate'])}",
                    f"وضعیت: {status}",
                    f"علت غیرفعال‌سازی: {inactive_reason}",
                ]
            )
        )
        blur = QGraphicsBlurEffect(self.stack)
        blur.setBlurRadius(7)
        self.stack.setGraphicsEffect(blur)
        self.profile_overlay.setGeometry(self.rect())
        self.profile_overlay.raise_()
        self.profile_overlay.show()

    def close_profile_overlay(self) -> None:
        self.stack.setGraphicsEffect(None)
        self.profile_overlay.hide()

    def resizeEvent(self, event) -> None:
        super().resizeEvent(event)
        if hasattr(self, "profile_overlay"):
            self.profile_overlay.setGeometry(self.rect())

    def clear_driver_checks(self) -> None:
        self._updating_checks = True
        for row in range(self.table.rowCount()):
            checkbox = self._checkbox_at_row(row)
            if checkbox:
                checkbox.setChecked(False)
        self._updating_checks = False

    def _checkbox_at_row(self, row: int) -> QCheckBox | None:
        holder = self.table.cellWidget(row, 0)
        if holder is None:
            return None
        return holder.findChild(QCheckBox)

    def _leave_report_selection_mode(self, clear_checks: bool = False) -> None:
        if not self.report_selection_mode:
            return
        self.report_selection_mode = False
        self.report_button.setText("گزارش از لیست")
        if clear_checks:
            self.clear_driver_checks()

    def _refresh_active_count(self) -> None:
        drivers = self.db.list_drivers()
        total = len(drivers)
        active_count = sum(1 for driver in drivers if int(driver.get("is_active", 1)))
        inactive_count = total - active_count
        self.total_count_label.setText(to_persian_digits(total))
        self.active_count_label.setText(to_persian_digits(active_count))
        self.inactive_count_label.setText(to_persian_digits(inactive_count))

    def clear_form(self) -> None:
        for line_edit in [
            self.first_name_input,
            self.last_name_input,
            self.mobile_input,
            self.national_id_input,
            self.birth_date_input,
            self.car_model_input,
            self.car_year_input,
            self.car_color_input,
            self.distance_rate_input,
        ]:
            line_edit.clear()
        self.vehicle_status_combo.setCurrentIndex(0)
        self._update_distance_rate_visibility()

    def _format_birth_date(self, text: str) -> None:
        if self._formatting_birth_date:
            return
        self._formatting_birth_date = True
        digits = "".join(ch for ch in to_english_digits(text) if ch.isdigit())[:8]
        if len(digits) <= 4:
            formatted = digits
        elif len(digits) <= 6:
            formatted = f"{digits[:4]}/{digits[4:]}"
        else:
            formatted = f"{digits[:4]}/{digits[4:6]}/{digits[6:]}"
        self.birth_date_input.setText(to_persian_digits(formatted))
        self.birth_date_input.setCursorPosition(len(formatted))
        self._ensure_input_alignment(self.birth_date_input)
        self._formatting_birth_date = False

    def _format_distance_rate(self, text: str) -> None:
        if self._formatting_distance_rate:
            return
        self._formatting_distance_rate = True
        number = self._plain_number(text)
        formatted = self._format_number(number) if number else ""
        self.distance_rate_input.setText(formatted)
        self.distance_rate_input.setCursorPosition(len(formatted))
        self._ensure_input_alignment(self.distance_rate_input)
        self._formatting_distance_rate = False

    def _driver_by_id(self, driver_id: int) -> dict | None:
        for driver in self.db.list_drivers():
            if int(driver["id"]) == driver_id:
                return driver
        return None

    @staticmethod
    def _is_persian_text(value: str) -> bool:
        return bool(value and PERSIAN_TEXT_RE.fullmatch(value)) and not any(ch.isdigit() for ch in value)

    @staticmethod
    def _is_valid_date(value: str) -> bool:
        if not DATE_RE.fullmatch(value):
            return False
        _, month, day = (int(part) for part in value.split("/"))
        return 1 <= month <= 12 and 1 <= day <= 31

    @staticmethod
    def _plain_number(value: int | float | str) -> int:
        if isinstance(value, (int, float)):
            return int(value)
        digits = "".join(ch for ch in str(value) if ch.isdigit())
        return int(digits) if digits else 0

    @staticmethod
    def _normalize_digits(value: str) -> str:
        translation = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
        return to_english_digits(value)

    @staticmethod
    def _format_number(value: int | float | str) -> str:
        number = DriversPage._plain_number(value)
        return to_persian_digits(f"{number:,}") if number else ""

    @staticmethod
    def _format_rial(value: int | float | str) -> str:
        number = DriversPage._plain_number(value)
        return f"{to_persian_digits(f'{number:,}')} ریال" if number else "۰ ریال"
