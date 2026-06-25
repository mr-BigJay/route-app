from __future__ import annotations

from pathlib import Path

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QComboBox,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QPushButton,
    QSizePolicy,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from database.db import DatabaseManager
from ui.utils import (
    PRIMARY_COLOR,
    SUCCESS_COLOR,
    Page,
    gregorian_to_jalali,
    make_stat_card,
    show_error,
    show_success,
    to_english_digits,
    to_persian_digits,
)


class ReportsPage(Page):
    def __init__(self, db: DatabaseManager) -> None:
        super().__init__("گزارشات", "گزارش‌گیری ساده و خروجی Excel")
        self.setObjectName("reportsPage")
        self.db = db
        self.report_rows: list[dict] = []
        self.total_distance = 0.0
        self.total_fee = 0.0
        self.root_layout.addWidget(self._filters_card())
        self.root_layout.addWidget(self._table_card(), stretch=1)
        self.root_layout.addWidget(self._totals_card())
        self._center_page_header()

    def _center_page_header(self) -> None:
        header = self.root_layout.itemAt(0).widget()
        if header is None:
            return
        layout = header.layout()
        if layout is None:
            return
        for index in range(layout.count()):
            item = layout.itemAt(index)
            widget = item.widget()
            if widget is not None:
                widget.setAlignment(Qt.AlignmentFlag.AlignCenter)
                layout.setAlignment(widget, Qt.AlignmentFlag.AlignHCenter)

    def _build_themed_card(self, title: str, subtitle: str | None = None) -> tuple[QFrame, QVBoxLayout]:
        card = self.card()
        card.setObjectName("missionsTableCard")
        layout = QVBoxLayout(card)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)

        header = QFrame()
        header.setObjectName("missionsTableHeader")
        header_layout = QVBoxLayout(header)
        header_layout.setContentsMargins(18, 10, 18, 10)
        header_layout.setSpacing(0)
        title_label = QLabel(title)
        title_label.setObjectName("missionsTableTitle")
        title_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        title_label.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        header_layout.addWidget(title_label)
        if subtitle:
            header_layout.setSpacing(4)
            header_layout.setContentsMargins(18, 14, 18, 12)
            subtitle_label = QLabel(subtitle)
            subtitle_label.setObjectName("missionsTableSubtitle")
            subtitle_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
            subtitle_label.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
            header_layout.addWidget(subtitle_label)

        wrap = QFrame()
        wrap.setObjectName("missionsTableWrap")
        body_layout = QVBoxLayout(wrap)
        body_layout.setContentsMargins(14, 14, 14, 14)
        body_layout.setSpacing(12)

        layout.addWidget(header)
        layout.addWidget(wrap, stretch=1)
        return card, body_layout

    def _style_filter_input(self, widget: QWidget) -> None:
        widget.setObjectName("reportsFilterInput")
        widget.setLayoutDirection(Qt.LayoutDirection.RightToLeft)

    def _filters_card(self) -> QFrame:
        card, layout = self._build_themed_card("گزارش ماموریت رانندگان")

        self.driver_combo = QComboBox()
        self._style_filter_input(self.driver_combo)
        self.start_date_input = QLineEdit()
        self.end_date_input = QLineEdit()
        for date_input in [self.start_date_input, self.end_date_input]:
            self._style_filter_input(date_input)
            date_input.setPlaceholderText("yyyy/mm/dd")
            date_input.setMaxLength(10)
            date_input.setAlignment(Qt.AlignmentFlag.AlignCenter)

        filters_container = QWidget()
        filters_container.setLayoutDirection(Qt.LayoutDirection.LeftToRight)
        filters_row = QHBoxLayout(filters_container)
        filters_row.setContentsMargins(0, 0, 0, 0)
        filters_row.setSpacing(12)
        filters_row.addWidget(self._inline_field("تاریخ انتها", self.end_date_input), stretch=1)
        filters_row.addWidget(self._inline_field("تاریخ ابتدا", self.start_date_input), stretch=1)
        filters_row.addWidget(self._inline_field("راننده", self.driver_combo), stretch=3)

        buttons = QHBoxLayout()
        buttons.setDirection(QHBoxLayout.Direction.RightToLeft)
        buttons.setSpacing(10)
        generate_button = self._reports_action_button("نمایش گزارش", "primary")
        export_button = self._reports_action_button("خروجی Excel", "secondary")
        generate_button.clicked.connect(self.generate_report)
        export_button.clicked.connect(self.export_excel)
        buttons.addWidget(generate_button)
        buttons.addWidget(export_button)
        buttons.addStretch(1)

        layout.addWidget(filters_container)
        layout.addLayout(buttons)
        return card

    def _reports_action_button(self, title: str, variant: str) -> QPushButton:
        button = QPushButton(title)
        button.setObjectName("reportsActionButton")
        button.setProperty("variant", variant)
        button.setCursor(Qt.CursorShape.PointingHandCursor)
        return button

    def _inline_field(self, label: str, widget: QWidget) -> QWidget:
        box = QWidget()
        box.setObjectName("reportsFilterField")
        box.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        layout = QHBoxLayout(box)
        layout.setDirection(QHBoxLayout.Direction.RightToLeft)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(6)
        label_widget = QLabel(label)
        label_widget.setObjectName("reportsFieldLabel")
        label_widget.setAlignment(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter)
        label_widget.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)
        widget.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        layout.addWidget(label_widget)
        layout.addWidget(widget, stretch=1)
        return box

    def _table_card(self) -> QFrame:
        card, layout = self._build_themed_card("نتایج گزارش")
        self.table = QTableWidget(0, 9)
        self.table.setObjectName("missionsTable")
        self.table.setHorizontalHeaderLabels(
            ["ردیف", "تاریخ", "ساعت", "راننده", "خودرو", "مبدا", "مقصد نهایی", "مسافت", "حق‌الزحمه"]
        )
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.horizontalHeader().setDefaultAlignment(Qt.AlignmentFlag.AlignCenter)
        self.table.verticalHeader().setVisible(False)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.setShowGrid(False)
        self.table.setFocusPolicy(Qt.FocusPolicy.StrongFocus)
        layout.addWidget(self.table, stretch=1)
        return card

    def _totals_card(self) -> QFrame:
        card = self.card()
        layout = QHBoxLayout(card)
        layout.setContentsMargins(18, 16, 18, 16)
        layout.setSpacing(12)
        self.distance_total_card = make_stat_card("مسافت کل", "۰.۰", PRIMARY_COLOR)
        self.fee_total_card = make_stat_card("حق‌الزحمه راننده", "۰ ریال", SUCCESS_COLOR)
        layout.addWidget(self.distance_total_card)
        layout.addWidget(self.fee_total_card)
        return card

    def refresh(self) -> None:
        self._refresh_driver_combo()
        self._set_default_previous_month()
        self.generate_report()

    def _refresh_driver_combo(self) -> None:
        current = self.driver_combo.currentData()
        self.driver_combo.clear()
        self.driver_combo.addItem("همه راننده‌ها", None)
        for driver in self.db.list_drivers():
            self.driver_combo.addItem(driver["full_name"], driver["id"])
        index = self.driver_combo.findData(current)
        if index >= 0:
            self.driver_combo.setCurrentIndex(index)

    def _set_default_previous_month(self) -> None:
        today = gregorian_to_jalali()
        year, month, _ = [int(part) for part in today.split("/")]
        month -= 1
        if month == 0:
            year -= 1
            month = 12
        last_day = 31 if month <= 6 else 30
        if month == 12:
            last_day = 29
        self.start_date_input.setText(to_persian_digits(f"{year:04d}/{month:02d}/01"))
        self.end_date_input.setText(to_persian_digits(f"{year:04d}/{month:02d}/{last_day:02d}"))

    def generate_report(self) -> None:
        start_date = to_english_digits(self.start_date_input.text().strip())
        end_date = to_english_digits(self.end_date_input.text().strip())
        if not self._valid_date(start_date) or not self._valid_date(end_date):
            show_error(self, "بازه تاریخ را با فرمت yyyy/mm/dd وارد کنید.")
            return

        driver_id = self.driver_combo.currentData()
        rows = self.db.list_missions(filters={"driver_id": driver_id} if driver_id else {})
        self.report_rows = [
            self._report_row(row)
            for row in rows
            if start_date <= row["mission_date"] <= end_date
        ]
        self.total_distance = sum(float(row["distance"] or 0) for row in self.report_rows)
        self.total_fee = sum(float(row["driver_fee"] or 0) for row in self.report_rows)
        self._render_table()
        self._render_totals()

    def _report_row(self, mission: dict) -> dict:
        distance = float(mission["distance"] or 0)
        rate = float(mission.get("driver_distance_rate") or 0)
        return {
            **mission,
            "final_destination": self._final_destination(mission["destination"]),
            "driver_fee": distance * rate,
        }

    def _render_table(self) -> None:
        self.table.setRowCount(len(self.report_rows))
        for row, mission in enumerate(self.report_rows):
            values = [
                to_persian_digits(row + 1),
                to_persian_digits(mission["mission_date"]),
                to_persian_digits(mission["mission_time"]),
                mission["driver_name"],
                mission["vehicle"],
                mission["origin"],
                mission["final_destination"],
                to_persian_digits(f"{float(mission['distance']):.1f}"),
                self._format_rial(mission["driver_fee"]),
            ]
            for col, value in enumerate(values):
                item = QTableWidgetItem(str(value))
                item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                self.table.setItem(row, col, item)

    def _render_totals(self) -> None:
        self.distance_total_card.deleteLater()
        self.fee_total_card.deleteLater()
        parent_layout = self.root_layout.itemAt(self.root_layout.count() - 1).widget().layout()
        self.distance_total_card = make_stat_card("مسافت کل", to_persian_digits(f"{self.total_distance:.1f}"), PRIMARY_COLOR)
        self.fee_total_card = make_stat_card("حق‌الزحمه راننده", self._format_rial(self.total_fee), SUCCESS_COLOR)
        parent_layout.addWidget(self.distance_total_card)
        parent_layout.addWidget(self.fee_total_card)

    def export_excel(self) -> None:
        if not self.report_rows:
            show_error(self, "داده‌ای برای خروجی Excel وجود ندارد.")
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            show_error(self, "کتابخانه openpyxl نصب نیست.")
            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "ذخیره گزارش",
            str(Path.home() / "route-report.xlsx"),
            "Excel Files (*.xlsx)",
        )
        if not file_path:
            return

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Route Report"
        sheet.sheet_view.rightToLeft = True
        headers = ["ردیف", "تاریخ", "ساعت", "راننده", "خودرو", "مبدا", "مقصد نهایی", "مسافت", "حق‌الزحمه"]
        sheet.append(headers)
        for index, mission in enumerate(self.report_rows, start=1):
            sheet.append(
                [
                    index,
                    mission["mission_date"],
                    mission["mission_time"],
                    mission["driver_name"],
                    mission["vehicle"],
                    mission["origin"],
                    mission["final_destination"],
                    float(mission["distance"] or 0),
                    float(mission["driver_fee"] or 0),
                ]
            )
        sheet.append([])
        sheet.append(["", "", "", "", "", "", "جمع", self.total_distance, self.total_fee])

        header_fill = PatternFill("solid", fgColor="2563EB")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[column[0].column_letter].width = min(max_length + 4, 45)
        workbook.save(file_path)
        show_success(self, "فایل Excel با موفقیت ذخیره شد.")

    @staticmethod
    def _final_destination(destination: str) -> str:
        parts = [part.strip() for part in destination.replace(",", "،").split("،") if part.strip()]
        return parts[-1] if parts else destination

    @staticmethod
    def _valid_date(value: str) -> bool:
        parts = value.split("/")
        if len(parts) != 3 or not all(part.isdigit() for part in parts):
            return False
        year, month, day = [int(part) for part in parts]
        return 1200 <= year <= 1600 and 1 <= month <= 12 and 1 <= day <= 31

    @staticmethod
    def _format_rial(value: float) -> str:
        return f"{to_persian_digits(f'{int(value):,}')} ریال"
